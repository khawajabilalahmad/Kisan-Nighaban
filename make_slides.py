from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[0] # 0 is title slide layout
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    title_shape.text = title
    subtitle_shape.text = subtitle
    return slide

def add_bullet_slide(prs, title, bullets):
    slide_layout = prs.slide_layouts[1] # 1 is title and content layout
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = title
    
    tf = body_shape.text_frame
    if bullets:
        tf.text = bullets[0]
        for bullet in bullets[1:]:
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
        
    return slide

prs = Presentation()

# Slide 1
add_title_slide(prs, "Kisan Nighaban (کسان نگہبان)", "Empowering Farmers with AI-Driven Climate & Crop Intelligence\n\nBuilt for the Competition")

# Slide 2
add_bullet_slide(prs, "The Challenges Facing Modern Agriculture", [
    "Unpredictable Climate: Sudden weather shifts destroy yields without warning.",
    "Lack of Instant Expertise: Farmers often rely on outdated, generalized advice rather than data-driven insights.",
    "The Language Barrier: Cutting-edge agritech tools are rarely available in local languages (Urdu/Hinglish).",
    "Fragmented Data: Tracking multiple farms, sowing dates, and crop types is done on paper or not at all."
])

# Slide 3
add_bullet_slide(prs, "Meet Kisan Nighaban - Your Digital 'Kisaan Bhai'", [
    "Kisan Nighaban is an intelligent, bilingual Progressive Web App (PWA).",
    "Acts as a 24/7 digital assistant for farmers.",
    "Combines real-time weather data, farm-specific tracking, and generative AI.",
    "Designed to protect crops and maximize yields."
])

# Slide 4
add_bullet_slide(prs, "Feature 1 - Intelligent Farm Management", [
    "Multiple Profiles: Farmers can monitor multiple fields simultaneously.",
    "Smart Tracking: Logs crop types, soil conditions, water sources, and precise sowing dates.",
    "Activity Logs: Keep a digital record of fertilizers, watering, and treatments to feed into the AI's memory."
])

# Slide 5
add_bullet_slide(prs, "Feature 2 - AI Climate & Risk Analysis", [
    "Real-Time Data Synthesis: Merges hyper-local weather forecasts with the farm's specific crop data.",
    "The Health Score: The AI calculates a dynamic 0-100 Health Score for the farm based on current conditions.",
    "Actionable Mitigation: Generates immediate, localized recommendations to prevent weather-related crop damage before it happens."
])

# Slide 6
add_bullet_slide(prs, "Feature 3 - Context-Aware Chatbot", [
    "Not Just a Chatbot: The AI is injected with the exact context of the user's farm.",
    "Knows crop age, soil type, latest weather, and recent activities.",
    "Personalized Solutions: Farmers can ask specific questions like, 'Why are my cotton leaves turning yellow?'",
    "Receives answers tailored to their exact farm conditions."
])

# Slide 7
add_bullet_slide(prs, "Feature 4 - Built for Everyone (Localization)", [
    "Native UI Localization: The entire app interface dynamically switches between English, Roman Urdu, and standard Urdu.",
    "Bilingual AI Brain: The Chatbot automatically detects the user's language preference.",
    "Translates its complex agricultural advice into perfectly fluent local languages."
])

# Slide 8
add_bullet_slide(prs, "UI/UX & Accessibility", [
    "Immersive Design: Features dynamic, time-based diorama backgrounds (day, evening, night) that make the app feel alive.",
    "Progressive Web App (PWA): Can be installed directly from the browser or packaged as an ultra-lightweight APK.",
    "Performance: Designed to run smoothly even on low-end Android devices in rural areas."
])

# Slide 9
add_bullet_slide(prs, "The Technology Stack", [
    "Frontend: React.js, Tailwind CSS (Glassmorphism), Vite, React Router, i18next.",
    "Backend: Python, FastAPI, SQLAlchemy, JWT Authentication.",
    "Intelligence Core: Google Gemini AI Model (Prompt Engineering & Context Injection).",
    "Deployment: Vercel (Frontend) & Render (Backend)."
])

# Slide 10
add_bullet_slide(prs, "What's Next for Kisan Nighaban?", [
    "IoT Integration: Direct connection to cheap soil moisture and pH sensors for automated alerts.",
    "Market Intelligence: Live crop pricing and marketplace connections to help farmers sell at the best price.",
    "Community Hub: A social feature for farmers in similar districts to share local insights."
])

# Slide 11
add_title_slide(prs, "Thank You!", "Protecting our crops today. Securing our food for tomorrow.\n\nTeam: Your Team Name")

prs.save('Kisan_Nighaban_Pitch.pptx')
print("Presentation generated successfully at Kisan_Nighaban_Pitch.pptx")
